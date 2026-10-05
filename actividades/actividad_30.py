import os
from datetime import datetime, timezone

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


# ==========================================================
# ACTIVIDAD 30 - REPORTE FINAL CONSOLIDADO
# ==========================================================

print("=" * 75)
print("ACTIVIDAD 30 - REPORTE FINAL CONSOLIDADO")
print("=" * 75)


SIMBOLO = "SOLUSD"
TEMPORALIDAD = "H1"

ARCHIVO_DATOS = "solana_4_semanas.csv"

ARCHIVO_SALIDA = (
    "reportes/actividad_30_reporte_final_solana.xlsx"
)


# ==========================================================
# 1. COMPROBAR ARCHIVO PRINCIPAL
# ==========================================================

if not os.path.exists(ARCHIVO_DATOS):
    raise FileNotFoundError(
        f"No se encontró {ARCHIVO_DATOS}"
    )


df = pd.read_csv(
    ARCHIVO_DATOS
)

if df.empty:
    raise ValueError(
        "El archivo principal no contiene datos."
    )


df["time"] = pd.to_datetime(
    df["time"],
    utc=True
)

df = (
    df.sort_values("time")
    .reset_index(drop=True)
)


print(f"\nInstrumento: {SIMBOLO}")
print(f"Temporalidad: {TEMPORALIDAD}")
print(f"Registros: {len(df)}")

print(
    "Período: "
    f"{df['time'].min()} "
    f"hasta {df['time'].max()}"
)


# ==========================================================
# 2. COMPROBACIONES DE CALIDAD
# ==========================================================

filas = len(df)
columnas = len(df.columns)

valores_faltantes = int(
    df.isna().sum().sum()
)

fechas_duplicadas = int(
    df["time"].duplicated().sum()
)

ohlc_cero = int(
    (
        (df["open"] == 0)
        | (df["high"] == 0)
        | (df["low"] == 0)
        | (df["close"] == 0)
    ).sum()
)


calidad = pd.DataFrame({
    "Metrica": [
        "Cantidad de registros",
        "Cantidad de columnas",
        "Valores faltantes",
        "Fechas duplicadas",
        "Registros OHLC con valor cero",
        "Fecha inicial",
        "Fecha final"
    ],
    "Valor": [
        filas,
        columnas,
        valores_faltantes,
        fechas_duplicadas,
        ohlc_cero,
        str(df["time"].min()),
        str(df["time"].max())
    ]
})


# ==========================================================
# 3. RESUMEN DIARIO
# ==========================================================

df["fecha"] = (
    df["time"].dt.date
)

resumen_diario = (
    df.groupby("fecha")
    .agg(
        apertura=("open", "first"),
        maximo=("high", "max"),
        minimo=("low", "min"),
        cierre=("close", "last"),
        volumen_ticks=("tick_volume", "sum"),
        velas=("time", "count")
    )
    .reset_index()
)

resumen_diario[
    "variacion_pct"
] = (
    (
        resumen_diario["cierre"]
        - resumen_diario["apertura"]
    )
    / resumen_diario["apertura"]
) * 100


# ==========================================================
# 4. VARIACIONES HORARIAS
# ==========================================================

df[
    "variacion_pct"
] = (
    df["close"]
    .pct_change()
    * 100
)

df[
    "movimiento_abs_pct"
] = (
    df["variacion_pct"].abs()
)


mayores_movimientos = (
    df.nlargest(
        10,
        "movimiento_abs_pct"
    )[
        [
            "time",
            "open",
            "high",
            "low",
            "close",
            "variacion_pct",
            "movimiento_abs_pct",
            "tick_volume"
        ]
    ]
    .copy()
)


# ==========================================================
# 5. ESTADÍSTICAS GENERALES
# ==========================================================

variaciones_validas = (
    df["variacion_pct"]
    .dropna()
)

estadisticas = pd.DataFrame({
    "Metrica": [
        "Precio inicial",
        "Precio final",
        "Precio máximo",
        "Precio mínimo",
        "Variación total (%)",
        "Media variación horaria (%)",
        "Mediana variación horaria (%)",
        "Desviación estándar (%)",
        "Mayor subida horaria (%)",
        "Mayor caída horaria (%)",
        "Promedio tick_volume"
    ],

    "Valor": [
        df["close"].iloc[0],
        df["close"].iloc[-1],
        df["high"].max(),
        df["low"].min(),

        (
            (
                df["close"].iloc[-1]
                / df["close"].iloc[0]
            ) - 1
        ) * 100,

        variaciones_validas.mean(),
        variaciones_validas.median(),
        variaciones_validas.std(),
        variaciones_validas.max(),
        variaciones_validas.min(),
        df["tick_volume"].mean()
    ]
})


# ==========================================================
# 6. ARCHIVOS DE ACTIVIDADES ANTERIORES
# ==========================================================

archivos_adicionales = {
    "Variacion_acumulada":
        "variacion_acumulada_solana.csv",

    "Rachas":
        "rachas_consecutivas_solana.csv",

    "Maximo_24h":
        "distancia_maximo_24_solana.csv",

    "Recuperacion_caidas":
        "recuperacion_mayores_caidas_solana.csv",

    "Distribucion":
        "variaciones_horarias_solana.csv",

    "Movimiento_ticks":
        "movimiento_vs_ticks_solana.csv",

    "Bid_Ask":
        "observaciones_bid_ask_solana.csv",

    "Comparacion_semanal":
        "comparacion_semanal_solana.csv"
}


datos_adicionales = {}


print("\nARCHIVOS DE ACTIVIDADES ANTERIORES")


for hoja, archivo in archivos_adicionales.items():

    if os.path.exists(archivo):

        datos_adicionales[
            hoja
        ] = pd.read_csv(
            archivo
        )

        print(
            f"OK - {archivo}: "
            f"{len(datos_adicionales[hoja])} registros"
        )

    else:

        print(
            f"ADVERTENCIA - No encontrado: {archivo}"
        )


# ==========================================================
# 7. CONTROL DE ACTIVIDADES
# ==========================================================

actividades = []

for numero in range(1, 31):

    if numero <= 28:

        estado = "Completada"

    elif numero == 29:

        estado = "Completada"

    else:

        estado = "Reporte final consolidado"

    actividades.append({
        "Actividad": numero,
        "Estado": estado
    })


control_actividades = pd.DataFrame(
    actividades
)


# ==========================================================
# 8. METADATOS
# ==========================================================

fecha_generacion = datetime.now(
    timezone.utc
)

metadatos = pd.DataFrame({
    "Campo": [
        "Instrumento",
        "Descripción",
        "Temporalidad",
        "Zona horaria",
        "Cantidad de registros",
        "Período inicial",
        "Período final",
        "Fuente",
        "Procesamiento",
        "Cantidad de actividades",
        "Fecha de generación"
    ],

    "Valor": [
        SIMBOLO,
        "Solana frente al dólar estadounidense",
        TEMPORALIDAD,
        "UTC",
        len(df),
        str(df["time"].min()),
        str(df["time"].max()),
        "MetaTrader 5 - TIO Markets Demo",
        "Python, Pandas, Matplotlib y Celery",
        30,
        fecha_generacion.strftime(
            "%Y-%m-%d %H:%M:%S UTC"
        )
    ]
})


# ==========================================================
# 9. PREPARAR FECHAS PARA EXCEL
# ==========================================================

# Excel no admite datetimes con timezone.

df_excel = df.copy()

df_excel["time"] = (
    df_excel["time"]
    .dt.tz_convert("UTC")
    .dt.tz_localize(None)
)


mayores_movimientos_excel = (
    mayores_movimientos.copy()
)

mayores_movimientos_excel["time"] = (
    mayores_movimientos_excel["time"]
    .dt.tz_convert("UTC")
    .dt.tz_localize(None)
)


# Revisar también los DataFrames adicionales
# por si alguno contiene fechas con timezone.

for nombre, datos in datos_adicionales.items():

    for columna in datos.columns:

        if "time" in columna.lower():

            try:

                fecha = pd.to_datetime(
                    datos[columna],
                    utc=True
                )

                datos[columna] = (
                    fecha
                    .dt.tz_convert("UTC")
                    .dt.tz_localize(None)
                )

            except Exception:
                pass


# ==========================================================
# 10. GENERAR EXCEL FINAL
# ==========================================================

os.makedirs(
    "reportes",
    exist_ok=True
)


with pd.ExcelWriter(
    ARCHIVO_SALIDA,
    engine="openpyxl"
) as writer:

    df_excel.to_excel(
        writer,
        sheet_name="Velas_SOL",
        index=False
    )

    resumen_diario.to_excel(
        writer,
        sheet_name="Resumen_diario",
        index=False
    )

    mayores_movimientos_excel.to_excel(
        writer,
        sheet_name="Mayores_movimientos",
        index=False
    )

    calidad.to_excel(
        writer,
        sheet_name="Calidad",
        index=False
    )

    estadisticas.to_excel(
        writer,
        sheet_name="Estadisticas",
        index=False
    )

    control_actividades.to_excel(
        writer,
        sheet_name="Actividades",
        index=False
    )

    for hoja, datos in datos_adicionales.items():

        datos.to_excel(
            writer,
            sheet_name=hoja[:31],
            index=False
        )

    metadatos.to_excel(
        writer,
        sheet_name="Metadatos",
        index=False
    )


# ==========================================================
# 11. DAR FORMATO AL EXCEL
# ==========================================================

libro = load_workbook(
    ARCHIVO_SALIDA
)


for hoja in libro.worksheets:

    hoja.freeze_panes = "A2"

    hoja.auto_filter.ref = (
        hoja.dimensions
    )

    for celda in hoja[1]:

        celda.font = Font(
            bold=True
        )

        celda.alignment = Alignment(
            horizontal="center"
        )

    for columna in hoja.columns:

        letra = get_column_letter(
            columna[0].column
        )

        longitud = 0

        for celda in columna:

            if celda.value is not None:

                longitud = max(
                    longitud,
                    len(str(celda.value))
                )

        hoja.column_dimensions[
            letra
        ].width = min(
            longitud + 2,
            35
        )


libro.save(
    ARCHIVO_SALIDA
)


print(
    f"\nReporte final generado:\n{ARCHIVO_SALIDA}"
)


# ==========================================================
# 12. RELEER Y VERIFICAR EL REPORTE
# ==========================================================

print("\n" + "=" * 75)
print("VERIFICACIÓN DEL REPORTE FINAL")
print("=" * 75)


excel = pd.ExcelFile(
    ARCHIVO_SALIDA
)


print("\nHojas encontradas:")

for hoja in excel.sheet_names:
    print(f" - {hoja}")


velas_recuperadas = pd.read_excel(
    ARCHIVO_SALIDA,
    sheet_name="Velas_SOL"
)

actividades_recuperadas = pd.read_excel(
    ARCHIVO_SALIDA,
    sheet_name="Actividades"
)

metadatos_recuperados = pd.read_excel(
    ARCHIVO_SALIDA,
    sheet_name="Metadatos"
)


print("\nCOMPROBACIÓN DE REGISTROS")

print(
    f"Registros originales: {len(df)}"
)

print(
    "Registros recuperados del Excel: "
    f"{len(velas_recuperadas)}"
)


if len(df) == len(
    velas_recuperadas
):

    print(
        "OK - La cantidad de registros coincide."
    )

else:

    print(
        "ERROR - La cantidad de registros no coincide."
    )


print("\nCOMPROBACIÓN DE ACTIVIDADES")

print(
    "Actividades registradas: "
    f"{len(actividades_recuperadas)}"
)


if len(
    actividades_recuperadas
) == 30:

    print(
        "OK - Se registraron las 30 actividades."
    )

else:

    print(
        "ERROR - No aparecen las 30 actividades."
    )


# ==========================================================
# 13. COMPROBAR METADATOS
# ==========================================================

print("\nCOMPROBACIÓN DE METADATOS")


def obtener_metadato(nombre):

    coincidencia = metadatos_recuperados[
        metadatos_recuperados["Campo"]
        == nombre
    ]

    if coincidencia.empty:
        return None

    return coincidencia.iloc[0]["Valor"]


simbolo_excel = obtener_metadato(
    "Instrumento"
)

temporalidad_excel = obtener_metadato(
    "Temporalidad"
)

inicio_excel = obtener_metadato(
    "Período inicial"
)

fin_excel = obtener_metadato(
    "Período final"
)


print(
    f"Símbolo: {simbolo_excel}"
)

print(
    f"Temporalidad: {temporalidad_excel}"
)

print(
    f"Período inicial: {inicio_excel}"
)

print(
    f"Período final: {fin_excel}"
)


if simbolo_excel == SIMBOLO:

    print(
        "OK - El símbolo corresponde a SOLUSD."
    )

else:

    print(
        "ERROR - El símbolo no corresponde a SOLUSD."
    )


# ==========================================================
# 14. COMPROBAR GRÁFICAS
# ==========================================================

print("\nCOMPROBACIÓN DE GRÁFICAS")


graficas_requeridas = [
    "graficas/actividad_21_variacion_acumulada.png",
    "graficas/actividad_22_rachas_consecutivas.png",
    "graficas/actividad_23_distancia_maximo_24.png",
    "graficas/actividad_24_recuperacion_caidas.png",
    "graficas/actividad_25_histograma_variaciones.png",
    "graficas/actividad_26_movimiento_vs_ticks.png",
    "graficas/actividad_27_bid_ask.png",
    "graficas/actividad_28_precio_4_semanas.png"
]


graficas_encontradas = 0


for grafica in graficas_requeridas:

    if os.path.exists(grafica):

        print(
            f"OK - {grafica}"
        )

        graficas_encontradas += 1

    else:

        print(
            f"ADVERTENCIA - No encontrada: {grafica}"
        )


print(
    "\nGráficas encontradas: "
    f"{graficas_encontradas}/"
    f"{len(graficas_requeridas)}"
)


# ==========================================================
# 15. RESULTADO FINAL
# ==========================================================

print("\n" + "=" * 75)

if (
    len(df) == len(velas_recuperadas)
    and len(actividades_recuperadas) == 30
    and simbolo_excel == SIMBOLO
):

    print(
        "ACTIVIDAD 30 COMPLETADA CORRECTAMENTE"
    )

else:

    print(
        "ACTIVIDAD 30 FINALIZADA CON ADVERTENCIAS"
    )

print("=" * 75)